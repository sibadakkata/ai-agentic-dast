"""Excel report exporter using openpyxl.

Generates a multi-sheet .xlsx file with:
  Sheet 1: Summary — scan metadata, scanner sources, severity counts
  Sheet 2: All Findings — every triaged finding with all fields
  Sheet 3: AI vs Triage — side-by-side AI severity vs triage verdict
  Sheet 4: AI Raw — AI scanner raw findings

Color-coded severity cells, auto-width columns, filters enabled.
"""

from __future__ import annotations

import os
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


SEV_COLORS = {
    "Critical": "8B0000",
    "High": "DC2626",
    "Medium": "F59E0B",
    "Low": "3B82F6",
    "Info": "6B7280",
    "Not Exploitable": "22C55E",
}

SOURCE_COLORS = {
    "ai": "8B5CF6",
}

VERDICT_COLORS = {
    "TRUE_POSITIVE": "DC2626",
    "FALSE_POSITIVE": "22C55E",
    "MANUAL_REVIEW": "F59E0B",
    "NEEDS_VERIFICATION": "6B7280",
}

HEADER_FILL = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
THIN_BORDER = Border(
    left=Side(style="thin", color="D1D5DB"),
    right=Side(style="thin", color="D1D5DB"),
    top=Side(style="thin", color="D1D5DB"),
    bottom=Side(style="thin", color="D1D5DB"),
)


def _auto_width(ws, min_width=10, max_width=60):
    """Auto-size columns based on content."""
    for col_cells in ws.columns:
        lengths = []
        for cell in col_cells:
            if cell.value:
                lengths.append(len(str(cell.value)))
        if lengths:
            best = min(max(max(lengths), min_width), max_width)
            letter = get_column_letter(col_cells[0].column)
            ws.column_dimensions[letter].width = best + 2


def _write_header(ws, headers):
    """Write styled header row."""
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = THIN_BORDER


def _sev_fill(severity):
    """Get fill color for a severity value."""
    color = SEV_COLORS.get(severity, "F3F4F6")
    return PatternFill(start_color=color, end_color=color, fill_type="solid")


def _source_fill(source):
    color = SOURCE_COLORS.get(source, "F3F4F6")
    return PatternFill(start_color=color, end_color=color, fill_type="solid")


def generate_excel(
    scan_id: str,
    metadata: dict,
    triaged_findings: list[dict],
    raw_findings: list[dict],
    output_dir: str | None = None,
) -> str:
    """Generate an Excel report and return the file path.

    Args:
        scan_id: Scan identifier
        metadata: Scan metadata dict (target, model, scanner, duration, cost, etc.)
        triaged_findings: List of triaged finding dicts (from triage engine)
        raw_findings: List of raw finding dicts (from all scanners)
        output_dir: Directory to write the file (defaults to results/reports/)

    Returns:
        Absolute path to the generated .xlsx file
    """
    wb = Workbook()

    # --- Sheet 1: Summary ---
    ws_sum = wb.active
    ws_sum.title = "Summary"

    ws_sum["A1"] = "AI Agentic DAST — Scan Report"
    ws_sum["A1"].font = Font(bold=True, size=16)
    ws_sum.merge_cells("A1:D1")

    summary_data = [
        ("Scan ID", scan_id),
        ("Target", metadata.get("target", "")),
        ("Model", metadata.get("model", "")),
        ("Scanner", metadata.get("scanner", "ai")),
        ("Scan Mode", metadata.get("scan_mode", "")),
        ("Duration", f"{metadata.get('duration_seconds', 0):.0f}s" if metadata.get("duration_seconds") else "-"),
        ("LLM Cost", f"${metadata.get('cost_usd', 0):.4f}" if metadata.get("cost_usd") else "-"),
        ("Generated", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
        ("", ""),
        ("Total Findings", len(triaged_findings)),
        ("AI Findings", len(raw_findings)),
    ]

    for row, (label, value) in enumerate(summary_data, 3):
        ws_sum.cell(row=row, column=1, value=label).font = Font(bold=True)
        ws_sum.cell(row=row, column=2, value=str(value))

    sev_start = len(summary_data) + 4
    ws_sum.cell(row=sev_start, column=1, value="Severity Breakdown").font = Font(bold=True, size=12)
    sev_counts = {}
    for f in triaged_findings:
        s = f.get("final_severity", "Info")
        sev_counts[s] = sev_counts.get(s, 0) + 1
    for i, (sev, count) in enumerate(sorted(sev_counts.items(), key=lambda x: list(SEV_COLORS.keys()).index(x[0]) if x[0] in SEV_COLORS else 99)):
        r = sev_start + 1 + i
        cell_sev = ws_sum.cell(row=r, column=1, value=sev)
        cell_sev.fill = _sev_fill(sev)
        cell_sev.font = Font(color="FFFFFF", bold=True)
        ws_sum.cell(row=r, column=2, value=count)

    verdict_start = sev_start + len(sev_counts) + 3
    ws_sum.cell(row=verdict_start, column=1, value="Verdict Breakdown").font = Font(bold=True, size=12)
    verdict_counts = {}
    for f in triaged_findings:
        v = f.get("verdict", "NEEDS_VERIFICATION")
        verdict_counts[v] = verdict_counts.get(v, 0) + 1
    for i, (verdict, count) in enumerate(verdict_counts.items()):
        r = verdict_start + 1 + i
        ws_sum.cell(row=r, column=1, value=verdict.replace("_", " ").title())
        ws_sum.cell(row=r, column=2, value=count)

    _auto_width(ws_sum)

    # --- Sheet 2: All Findings ---
    ws_all = wb.create_sheet("All Findings")
    headers = [
        "#", "Title", "Source", "AI Severity", "Triage Severity", "Verdict",
        "CWE", "CVSS", "CVSS Rationale", "URL", "Verified", "Verification Method",
        "Reason", "Confidence"
    ]
    _write_header(ws_all, headers)

    for i, f in enumerate(triaged_findings, 1):
        row = i + 1
        ws_all.cell(row=row, column=1, value=i)
        ws_all.cell(row=row, column=2, value=f.get("title", ""))
        src_cell = ws_all.cell(row=row, column=3, value=(f.get("source", "ai")).upper())
        src_cell.fill = _source_fill(f.get("source", "ai"))
        src_cell.font = Font(color="FFFFFF", bold=True)
        ws_all.cell(row=row, column=4, value=f.get("ai_severity", ""))
        sev_cell = ws_all.cell(row=row, column=5, value=f.get("final_severity", ""))
        sev_cell.fill = _sev_fill(f.get("final_severity", ""))
        sev_cell.font = Font(color="FFFFFF", bold=True)
        verdict_cell = ws_all.cell(row=row, column=6, value=f.get("verdict", "").replace("_", " ").title())
        vc = VERDICT_COLORS.get(f.get("verdict", ""), "6B7280")
        verdict_cell.fill = PatternFill(start_color=vc, end_color=vc, fill_type="solid")
        verdict_cell.font = Font(color="FFFFFF", bold=True)
        ws_all.cell(row=row, column=7, value=f.get("cwe", ""))
        cvss_val = f.get("cvss_override") if f.get("cvss_override") is not None else f.get("cvss")
        ws_all.cell(row=row, column=8, value=round(cvss_val, 1) if cvss_val else 0)
        ws_all.cell(row=row, column=9, value=f.get("cvss_rationale", "")[:500])
        ws_all.cell(row=row, column=10, value=f.get("url", ""))
        ws_all.cell(row=row, column=11, value="Yes" if f.get("verified") else "No")
        ws_all.cell(row=row, column=12, value=f.get("verification_method", ""))
        ws_all.cell(row=row, column=13, value=f.get("reason", "")[:500])
        ws_all.cell(row=row, column=14, value=f.get("confidence_score", ""))

        for col in range(1, len(headers) + 1):
            ws_all.cell(row=row, column=col).border = THIN_BORDER

    ws_all.auto_filter.ref = ws_all.dimensions
    _auto_width(ws_all)

    # --- Sheet 3: AI vs Triage ---
    # Side-by-side comparison: what the AI said vs what triage decided.
    # Rows where severity differs between AI and triage are highlighted yellow
    # and sorted to the top so divergences jump out.
    ws_cmp = wb.create_sheet("AI vs Triage")
    cmp_headers = [
        "#", "Title", "URL", "Parameter",
        "AI Severity", "Triage Severity", "Changed?",
        "Verdict", "Verified", "Verification Method",
        "CWE", "CVSS", "Triage Reason",
    ]
    _write_header(ws_cmp, cmp_headers)

    changed_fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
    yes_fill = PatternFill(start_color="DC2626", end_color="DC2626", fill_type="solid")
    no_fill = PatternFill(start_color="22C55E", end_color="22C55E", fill_type="solid")

    def _cmp_key(f):
        ai = (f.get("ai_severity") or "").strip().lower()
        tr = (f.get("final_severity") or f.get("severity") or "").strip().lower()
        return (0 if ai != tr else 1, f.get("title", ""))

    sorted_cmp = sorted(triaged_findings, key=_cmp_key)
    changed_count = 0
    for i, f in enumerate(sorted_cmp, 1):
        row = i + 1
        ai_sev = (f.get("ai_severity") or "").strip()
        triage_sev = (f.get("final_severity") or f.get("severity") or "").strip()
        changed = ai_sev.lower() != triage_sev.lower()
        if changed:
            changed_count += 1

        ws_cmp.cell(row=row, column=1, value=i)
        ws_cmp.cell(row=row, column=2, value=f.get("title", ""))
        ws_cmp.cell(row=row, column=3, value=f.get("url", ""))
        ws_cmp.cell(row=row, column=4, value=f.get("parameter", ""))

        ai_cell = ws_cmp.cell(row=row, column=5, value=ai_sev or "-")
        if ai_sev:
            ai_cell.fill = _sev_fill(ai_sev)
            ai_cell.font = Font(color="FFFFFF", bold=True)

        tr_cell = ws_cmp.cell(row=row, column=6, value=triage_sev or "-")
        if triage_sev:
            tr_cell.fill = _sev_fill(triage_sev)
            tr_cell.font = Font(color="FFFFFF", bold=True)

        chg_cell = ws_cmp.cell(row=row, column=7, value="YES" if changed else "no")
        chg_cell.fill = yes_fill if changed else no_fill
        chg_cell.font = Font(color="FFFFFF", bold=True)
        chg_cell.alignment = Alignment(horizontal="center")

        verdict_cell = ws_cmp.cell(
            row=row, column=8,
            value=(f.get("verdict") or "").replace("_", " ").title(),
        )
        vc = VERDICT_COLORS.get(f.get("verdict", ""), "6B7280")
        verdict_cell.fill = PatternFill(start_color=vc, end_color=vc, fill_type="solid")
        verdict_cell.font = Font(color="FFFFFF", bold=True)

        ws_cmp.cell(row=row, column=9, value="Yes" if f.get("verified") else "No")
        ws_cmp.cell(row=row, column=10, value=f.get("verification_method", ""))
        ws_cmp.cell(row=row, column=11, value=f.get("cwe", ""))
        cvss_val = f.get("cvss_override") if f.get("cvss_override") is not None else f.get("cvss")
        ws_cmp.cell(row=row, column=12, value=round(cvss_val, 1) if cvss_val else 0)
        ws_cmp.cell(row=row, column=13, value=(f.get("reason") or "")[:500])

        if changed:
            # Soft-highlight the whole row to make divergences visible.
            for col in range(1, len(cmp_headers) + 1):
                cell = ws_cmp.cell(row=row, column=col)
                if cell.fill.start_color.rgb in (None, "00000000", "FFFFFFFF"):
                    cell.fill = changed_fill

        for col in range(1, len(cmp_headers) + 1):
            ws_cmp.cell(row=row, column=col).border = THIN_BORDER

    # Header note above the table explaining the sheet
    ws_cmp.insert_rows(1)
    note_cell = ws_cmp.cell(
        row=1, column=1,
        value=f"AI vs Triage — {changed_count} of {len(sorted_cmp)} finding(s) had severity changed by triage. "
              f"Rows where AI and triage disagreed are highlighted and sorted to the top.",
    )
    note_cell.font = Font(italic=True, color="6B7280")
    ws_cmp.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(cmp_headers))

    ws_cmp.auto_filter.ref = f"A2:{get_column_letter(len(cmp_headers))}{ws_cmp.max_row}"
    ws_cmp.freeze_panes = "A3"
    _auto_width(ws_cmp)

    # --- Sheet 4: AI Raw ---
    ws_ai = wb.create_sheet("AI Raw")
    ai_headers = ["#", "Title", "Severity", "URL", "Parameter", "Payload", "Evidence", "Confidence"]
    _write_header(ws_ai, ai_headers)
    ai_findings = [f for f in raw_findings if f.get("source", "ai") in ("ai", "both")]
    for i, f in enumerate(ai_findings, 1):
        row = i + 1
        ws_ai.cell(row=row, column=1, value=i)
        ws_ai.cell(row=row, column=2, value=f.get("title", ""))
        ws_ai.cell(row=row, column=3, value=f.get("severity", ""))
        ws_ai.cell(row=row, column=4, value=f.get("url", ""))
        ws_ai.cell(row=row, column=5, value=f.get("parameter", ""))
        ws_ai.cell(row=row, column=6, value=str(f.get("payload", ""))[:500])
        ws_ai.cell(row=row, column=7, value=str(f.get("evidence", ""))[:500])
        ws_ai.cell(row=row, column=8, value=str(f.get("confidence", "")))
    ws_ai.auto_filter.ref = ws_ai.dimensions
    _auto_width(ws_ai)

    # Save
    if not output_dir:
        output_dir = str(Path(__file__).resolve().parent.parent / "results" / "reports")
    os.makedirs(output_dir, exist_ok=True)
    filename = f"scan_report_{scan_id}.xlsx"
    filepath = os.path.join(output_dir, filename)
    wb.save(filepath)
    return filepath
