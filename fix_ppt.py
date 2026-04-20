from pptx import Presentation
from pptx.util import Inches, Pt, Emu
import copy

INPUT_FILE = r'C:\Temp\RedTeam_Git Action Security.pptx'
OUTPUT_FILE = r'C:\Temp\RedTeam_Git Action Security.pptx'

prs = Presentation(INPUT_FILE)

replacements = [
    # Slide 5: double space before Events
    ("GitHub Actions Workflow  Events", "GitHub Actions Workflow Events"),
    # Slide 5: extra space in (SCA )
    ("(SCA )", "(SCA)"),
    # Slide 5: Un-sanitized -> Unsanitized
    ("Un-sanitized", "Unsanitized"),
    # Slide 6: Step Security -> StepSecurity
    ("Step Security.", "StepSecurity."),
    # Slide 7: Harder Runner -> Harden-Runner
    ("Harder Runner", "Harden-Runner"),
    # Slide 8: spacing in ( 0 -10 )
    ("( 0 -10 )", "(0\u201310)"),
    # Slide 11: secyrit -> security
    ("secyrit team", "security team"),
    # Slide 11: breached -> breach, grammar fix
    ("security breached in Git Action if does happen",
     "security breach in Git Actions if one does happen"),
    # Slide 11: StepSecuty -> StepSecurity
    ("StepSecuty", "StepSecurity"),
    # Slide 11: secueity -> security
    ("secueity", "security"),
    # Slide 11: provide -> provides (grammar)
    ("provide E2E", "provides E2E"),
    # Slide 11: Proiviced -> Provide
    ("Proiviced required inputs", "Provide required inputs"),
]

fix_count = 0

for slide in prs.slides:
    for shape in slide.shapes:
        if shape.has_text_frame:
            for para in shape.text_frame.paragraphs:
                for run in para.runs:
                    original = run.text
                    new_text = original
                    for old, new in replacements:
                        if old in new_text:
                            new_text = new_text.replace(old, new)
                    if new_text != original:
                        print(f"  FIXED: \"{original}\"")
                        print(f"     ->  \"{new_text}\"")
                        run.text = new_text
                        fix_count += 1

print(f"\nTotal fixes applied: {fix_count}")
prs.save(OUTPUT_FILE)
print(f"Saved to: {OUTPUT_FILE}")
