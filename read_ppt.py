from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.enum.text import PP_ALIGN

prs = Presentation(r'C:\Temp\RedTeam_Git Action Security.pptx')

print(f'Slide width: {prs.slide_width}, height: {prs.slide_height}')
print(f'Number of slides: {len(prs.slides)}')
print(f'Slide layouts available: {len(prs.slide_layouts)}')
for i, layout in enumerate(prs.slide_layouts):
    print(f'  Layout {i}: {layout.name}')
print()

for idx, slide in enumerate(prs.slides):
    print(f'=== SLIDE {idx+1} ===')
    print(f'  Layout: {slide.slide_layout.name}')
    for shape in slide.shapes:
        print(f'  Shape: {shape.shape_type}, name={shape.name}, pos=({shape.left},{shape.top}), size=({shape.width},{shape.height})')
        if shape.has_text_frame:
            for pi, para in enumerate(shape.text_frame.paragraphs):
                text = para.text.strip()
                if text:
                    font_info = ''
                    if para.runs:
                        r = para.runs[0]
                        try:
                            color = r.font.color.rgb if r.font.color and r.font.color.rgb else "none"
                        except:
                            color = "theme"
                        font_info = f' [font={r.font.name}, size={r.font.size}, bold={r.font.bold}, color={color}]'
                    print(f'    P{pi}: "{text}"{font_info}')
        if shape.has_table:
            table = shape.table
            print(f'    TABLE: {len(table.rows)} rows x {len(table.columns)} cols')
            for ri, row in enumerate(table.rows):
                cells = [cell.text.strip() for cell in row.cells]
                print(f'      Row {ri}: {cells}')
    print()
